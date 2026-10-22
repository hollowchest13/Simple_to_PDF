import logging
import subprocess
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from simple_to_pdf.converters.lib_office_converter import LibreOfficeConverter
from simple_to_pdf.converters.models import ConversionResult


def make_result(success=(), failed=()) -> ConversionResult:
    res = ConversionResult()
    res.success.extend(success)
    res.failed.extend(failed)
    return res


@pytest.fixture
def converter() -> LibreOfficeConverter:
    """Экземпляр без вызова __init__ (он тянет базовые классы и LibreOffice)."""
    conv = LibreOfficeConverter.__new__(LibreOfficeConverter)
    conv.soffice_path = "/usr/bin/soffice"
    conv.chunk_size = 2
    conv.check_stop = MagicMock()
    return conv


def touch(path: Path, content: bytes = b"x") -> Path:
    path.write_bytes(content)
    return path


class TestConvertToPdf:
    def test_categorizes_files_and_aggregates_results(self, converter, tmp_path):
        xlsx = touch(tmp_path / "a.xlsx")
        docx = touch(tmp_path / "b.docx")
        pptx = touch(tmp_path / "c.pptx")
        png = touch(tmp_path / "d.png")
        unknown = touch(tmp_path / "e.bin")
        missing = tmp_path / "missing.docx"  # не создаём

        converter.is_table_file = lambda *, file_path: file_path.suffix == ".xlsx"
        converter.is_document_file = lambda *, file_path: (
            file_path.suffix.lower()
            in {
                ".docx",
                ".bin_doc",
            }
        )
        converter.is_presentation_file = lambda *, file_path: (
            file_path.suffix == ".pptx"
        )
        converter.is_image_file = lambda *, file_path: file_path.suffix == ".png"

        converter._convert_docs_to_pdf = MagicMock(
            return_value=make_result(
                success=[(0, tmp_path / "0.pdf"), (1, tmp_path / "1.pdf")],
                failed=[(2, pptx)],
            )
        )
        converter._convert_images_to_pdf = MagicMock(
            return_value=make_result(success=[(3, tmp_path / "3.pdf")])
        )

        files = [(0, xlsx), (1, docx), (2, pptx), (3, png), (4, unknown), (5, missing)]
        result = converter.convert_to_pdf(files=files, out_dir=tmp_path)

        converter._convert_docs_to_pdf.assert_called_once_with(
            files=[(0, xlsx), (1, docx), (2, pptx)], out_dir=tmp_path
        )
        converter._convert_images_to_pdf.assert_called_once_with(
            files=[(3, png)], out_dir=tmp_path
        )
        assert [i for i, _ in result.success] == [0, 1, 3]
        assert result.failed == [(2, pptx)]

    def test_empty_input(self, converter, tmp_path):
        converter._convert_docs_to_pdf = MagicMock(return_value=make_result())
        converter._convert_images_to_pdf = MagicMock(return_value=make_result())

        result = converter.convert_to_pdf(files=[], out_dir=tmp_path)

        assert result.success == []
        assert result.failed == []
        converter._convert_docs_to_pdf.assert_called_once_with(
            files=[], out_dir=tmp_path
        )
        converter._convert_images_to_pdf.assert_called_once_with(
            files=[], out_dir=tmp_path
        )


class TestConvertDocsToPdf:
    def test_processes_each_chunk_and_aggregates(self, converter, tmp_path):
        chunk1 = [(0, tmp_path / "a.docx"), (1, tmp_path / "b.docx")]
        chunk2 = [(2, tmp_path / "c.docx")]
        converter.make_chunks = MagicMock(return_value=[chunk1, chunk2])
        converter._convert_chunk = MagicMock(
            side_effect=[
                make_result(success=[(0, Path("0.pdf"))], failed=[(1, chunk1[1][1])]),
                make_result(success=[(2, Path("2.pdf"))]),
            ]
        )

        result = converter._convert_docs_to_pdf(files=chunk1 + chunk2, out_dir=tmp_path)

        converter.make_chunks.assert_called_once_with(
            chunk1 + chunk2, converter.chunk_size
        )
        assert converter.check_stop.call_count == 2
        assert converter._convert_chunk.call_count == 2
        assert result.success == [(0, Path("0.pdf")), (2, Path("2.pdf"))]
        assert result.failed == [(1, chunk1[1][1])]

    def test_stop_signal_propagates(self, converter, tmp_path):
        converter.make_chunks = MagicMock(return_value=[[(0, tmp_path / "a.docx")]])
        converter.check_stop = MagicMock(side_effect=RuntimeError("stopped"))
        converter._convert_chunk = MagicMock()

        with pytest.raises(RuntimeError, match="stopped"):
            converter._convert_docs_to_pdf(
                files=[(0, tmp_path / "a.docx")], out_dir=tmp_path
            )

        converter._convert_chunk.assert_not_called()

    def test_no_files(self, converter, tmp_path):
        converter.make_chunks = MagicMock(return_value=[])
        result = converter._convert_docs_to_pdf(files=[], out_dir=tmp_path)
        assert result.success == []
        assert result.failed == []


class TestRunFormatConversion:
    def test_success_builds_command_and_deletes_sources(self, converter, tmp_path):
        f1 = touch(tmp_path / "a.xls")
        f2 = touch(tmp_path / "b.csv")
        out_dir = tmp_path / "out"
        out_dir.mkdir()

        with patch.object(subprocess, "run") as run:
            converter._run_libreoffice_format_conversion(
                input_paths=[f1, f2], out_dir=out_dir
            )

        run.assert_called_once()
        command = run.call_args.args[0]
        assert command[:6] == [
            "/usr/bin/soffice",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
        ]
        assert command[6:] == [str(f1), str(f2)]
        assert run.call_args.kwargs["timeout"] == 10 + 2 * 5
        assert run.call_args.kwargs["check"] is True
        assert run.call_args.kwargs["capture_output"] is True
        assert not f1.exists()
        assert not f2.exists()

    def test_timeout_is_logged_and_sources_kept(self, converter, tmp_path, caplog):
        f1 = touch(tmp_path / "a.xls")

        with (
            patch.object(
                subprocess,
                "run",
                side_effect=subprocess.TimeoutExpired(cmd="soffice", timeout=15),
            ),
            caplog.at_level(logging.ERROR),
        ):
            converter._run_libreoffice_format_conversion(
                input_paths=[f1], out_dir=tmp_path
            )

        assert "timed out" in caplog.text
        assert f1.exists()

    def test_called_process_error_is_logged(self, converter, tmp_path, caplog):
        f1 = touch(tmp_path / "a.xls")
        error = subprocess.CalledProcessError(1, "soffice", stderr=b"boom")

        with (
            patch.object(subprocess, "run", side_effect=error),
            caplog.at_level(logging.ERROR),
        ):
            converter._run_libreoffice_format_conversion(
                input_paths=[f1], out_dir=tmp_path
            )

        assert "boom" in caplog.text
        assert f1.exists()

    def test_unexpected_error_does_not_raise(self, converter, tmp_path, caplog):
        f1 = touch(tmp_path / "a.xls")

        with (
            patch.object(subprocess, "run", side_effect=OSError("no soffice")),
            caplog.at_level(logging.ERROR),
        ):
            converter._run_libreoffice_format_conversion(
                input_paths=[f1], out_dir=tmp_path
            )

        assert "Unexpected error" in caplog.text
        assert f1.exists()


class TestUpdatePaths:
    EXTS = [".xls", ".xlsb", ".ods", ".csv"]  # noqa: RUF012

    def test_replaces_with_existing_xlsx(self, converter, tmp_path):
        xls = touch(tmp_path / "a.xls")
        xlsx = touch(tmp_path / "a.xlsx")

        result = converter._update_paths(all_paths=[xls], to_check_exts=self.EXTS)

        assert result == [xlsx]

    def test_keeps_original_and_warns_if_xlsx_missing(
        self, converter, tmp_path, caplog
    ):
        ods = touch(tmp_path / "a.ods")

        with caplog.at_level(logging.WARNING):
            result = converter._update_paths(all_paths=[ods], to_check_exts=self.EXTS)

        assert result == [ods]
        assert "a.ods" in caplog.text

    def test_other_extensions_untouched(self, converter, tmp_path):
        docx = touch(tmp_path / "a.docx")
        xlsx = touch(tmp_path / "b.xlsx")

        result = converter._update_paths(
            all_paths=[docx, xlsx], to_check_exts=self.EXTS
        )

        assert result == [docx, xlsx]

    def test_case_insensitive_extension(self, converter, tmp_path):
        upper = touch(tmp_path / "a.XLS")
        xlsx = touch(tmp_path / "a.xlsx")

        result = converter._update_paths(all_paths=[upper], to_check_exts=self.EXTS)

        assert result == [xlsx]

    def test_preserves_order_and_mixed_results(self, converter, tmp_path):
        a_xls = touch(tmp_path / "a.xls")
        touch(tmp_path / "a.xlsx")
        b_csv = touch(tmp_path / "b.csv")  # xlsx нет
        c_docx = touch(tmp_path / "c.docx")

        result = converter._update_paths(
            all_paths=[a_xls, b_csv, c_docx], to_check_exts=self.EXTS
        )

        assert result == [tmp_path / "a.xlsx", b_csv, c_docx]


def _fake_soffice(*pdf_names: str):
    """Simulates soffice: creates PDF files in the --outdir directory."""

    def _run(command, **kwargs):
        outdir = Path(command[command.index("--outdir") + 1])
        for name in pdf_names:
            (outdir / name).write_bytes(b"%PDF-1.4")
        return MagicMock(returncode=0)

    return _run


class TestRunLibreOfficeCommand:
    def test_success_moves_pdfs_to_out_dir(self, converter, tmp_path):
        out_dir = tmp_path / "out"
        out_dir.mkdir()

        with patch.object(
            subprocess, "run", side_effect=_fake_soffice("0_a.pdf", "1_b.pdf")
        ) as run:
            ok = converter._run_libreoffice_command(
                input_paths=["/tmp/0_a.docx", "/tmp/1_b.docx"], out_dir=out_dir
            )

        assert ok is True
        assert (out_dir / "0_a.pdf").exists()
        assert (out_dir / "1_b.pdf").exists()
        command = run.call_args.args[0]
        assert command[:4] == ["/usr/bin/soffice", "--headless", "--convert-to", "pdf"]
        assert command[-2:] == ["/tmp/0_a.docx", "/tmp/1_b.docx"]
        assert run.call_args.kwargs["timeout"] == 10 + 2 * 5

    def test_name_collision_does_not_overwrite(self, converter, tmp_path):
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        existing = touch(out_dir / "0_a.pdf", b"OLD")

        with patch.object(subprocess, "run", side_effect=_fake_soffice("0_a.pdf")):
            ok = converter._run_libreoffice_command(
                input_paths=["/tmp/0_a.docx"], out_dir=out_dir
            )

        assert ok is True
        assert existing.read_bytes() == b"OLD"
        assert len(list(out_dir.glob("*.pdf"))) == 2

    @pytest.mark.parametrize(
        "error",
        [
            subprocess.TimeoutExpired(cmd="soffice", timeout=1),
            subprocess.CalledProcessError(1, "soffice"),
            OSError("boom"),
        ],
        ids=["timeout", "called_process_error", "unexpected"],
    )
    def test_errors_return_false(self, converter, tmp_path, error):
        with patch.object(subprocess, "run", side_effect=error):
            ok = converter._run_libreoffice_command(
                input_paths=["/tmp/0_a.docx"], out_dir=tmp_path
            )

        assert ok is False
        assert list(tmp_path.glob("*.pdf")) == []


class TestCollectResults:
    def test_splits_success_and_failed(self, converter, tmp_path):
        ok_src = tmp_path / "src" / "a.docx"
        bad_src = tmp_path / "src" / "b.docx"
        pdf = touch(tmp_path / "0_a.pdf")

        res = converter._collect_results(
            chunk=[(0, ok_src), (1, bad_src)], tmp_dir=tmp_path
        )

        assert res.success == [(0, pdf)]
        assert res.failed == [(1, bad_src)]

    def test_logs_warning_on_missing_pdf(self, converter, tmp_path, caplog):
        with caplog.at_level(logging.WARNING):
            converter._collect_results(chunk=[(7, Path("x.docx"))], tmp_dir=tmp_path)

        assert "7_x.pdf" in caplog.text

    def test_empty_chunk(self, converter, tmp_path):
        res = converter._collect_results(chunk=[], tmp_dir=tmp_path)
        assert res.success == []
        assert res.failed == []


class TestConvertChunk:
    @pytest.fixture
    def chunk_env(self, converter, tmp_path):
        """Overrides all pipeline steps; _prepare_temp_files creates files in tmp_dir."""
        self_created: dict[str, list[Path]] = {}

        def prepare_temp_files(*, chunk, tmp_dir):
            paths = []
            for idx, p in chunk:
                new = tmp_dir / f"{idx}_{p.name}"
                new.write_bytes(b"x")
                paths.append(new)
            self_created["paths"] = paths
            return paths

        converter._prepare_temp_files = MagicMock(side_effect=prepare_temp_files)
        converter._run_libreoffice_format_conversion = MagicMock()
        converter._prepare_excel_scaling = MagicMock()
        converter._run_libreoffice_command = MagicMock(return_value=True)
        converter._collect_results = MagicMock(
            return_value=make_result(success=[(0, tmp_path / "0.pdf")])
        )
        converter._created = self_created
        return converter

    def test_plain_documents_skip_format_conversion(self, chunk_env, tmp_path):
        chunk = [(0, Path("a.docx")), (1, Path("b.pptx"))]

        res = chunk_env._convert_chunk(chunk=chunk, out_dir=tmp_path)

        chunk_env._run_libreoffice_format_conversion.assert_not_called()
        chunk_env._prepare_excel_scaling.assert_not_called()
        run_kwargs = chunk_env._run_libreoffice_command.call_args.kwargs
        assert [Path(p).name for p in run_kwargs["input_paths"]] == [
            "0_a.docx",
            "1_b.pptx",
        ]
        assert run_kwargs["out_dir"] == tmp_path
        chunk_env._collect_results.assert_called_once_with(
            chunk=chunk, tmp_dir=tmp_path
        )
        assert res.success == [(0, tmp_path / "0.pdf")]

    def test_legacy_tables_are_converted_to_xlsx_and_scaled(self, chunk_env, tmp_path):
        def fake_format_conversion(*, input_paths, out_dir):
            for p in input_paths:
                (out_dir / f"{p.stem}.xlsx").write_bytes(b"x")

        chunk_env._run_libreoffice_format_conversion.side_effect = (
            fake_format_conversion
        )
        chunk = [(0, Path("a.xls")), (1, Path("b.csv")), (2, Path("c.docx"))]

        chunk_env._convert_chunk(chunk=chunk, out_dir=tmp_path)

        converted = chunk_env._run_libreoffice_format_conversion.call_args.kwargs[
            "input_paths"
        ]
        assert sorted(p.suffix for p in converted) == [".csv", ".xls"]

        input_paths = chunk_env._run_libreoffice_command.call_args.kwargs["input_paths"]
        assert [Path(p).suffix for p in input_paths] == [".xlsx", ".xlsx", ".docx"]
        scaled = [
            c.kwargs["file_path"]
            for c in chunk_env._prepare_excel_scaling.call_args_list
        ]
        assert [p.suffix for p in scaled] == [".xlsx", ".xlsx"]

    def test_xlsx_is_scaled_without_format_conversion(self, chunk_env, tmp_path):
        chunk_env._convert_chunk(chunk=[(0, Path("a.xlsx"))], out_dir=tmp_path)

        chunk_env._run_libreoffice_format_conversion.assert_not_called()
        chunk_env._prepare_excel_scaling.assert_called_once()

    def test_failed_libreoffice_returns_empty_result(self, chunk_env, tmp_path):
        chunk_env._run_libreoffice_command.return_value = False

        res = chunk_env._convert_chunk(chunk=[(0, Path("a.docx"))], out_dir=tmp_path)

        chunk_env._collect_results.assert_not_called()
        assert res.success == []
        assert res.failed == []


class TestGetExcelWidth:
    def test_returns_max_column_per_sheet(self, converter, tmp_path):
        path = tmp_path / "book.xlsx"
        wb = openpyxl.Workbook()
        ws1 = wb.active
        assert ws1 is not None
        ws1.title = "Narrow"
        ws1.append(["a", "b"])
        ws2 = wb.create_sheet("Wide")
        ws2.append(list("abcdefg"))
        wb.save(path)

        report = converter.get_excel_width(file_path=path)

        assert report == {"Narrow": 2, "Wide": 7}

    def test_uses_data_only_and_closes_workbook(self, converter, tmp_path):
        fake_sheet = MagicMock(max_column=5)
        fake_wb = MagicMock()
        fake_wb.sheetnames = ["S1"]
        fake_wb.__getitem__.return_value = fake_sheet

        with patch("openpyxl.load_workbook", return_value=fake_wb) as load:
            report = converter.get_excel_width(file_path=tmp_path / "x.xlsx")

        load.assert_called_once_with(filename=tmp_path / "x.xlsx", data_only=True)
        fake_wb.close.assert_called_once()
        assert report == {"S1": 5}

    def test_zero_max_column_gives_zero(self, converter, tmp_path):
        fake_wb = MagicMock()
        fake_wb.sheetnames = ["Empty"]
        fake_wb.__getitem__.return_value = MagicMock(max_column=0)

        with patch("openpyxl.load_workbook", return_value=fake_wb):
            report = converter.get_excel_width(file_path=tmp_path / "x.xlsx")

        assert report == {"Empty": 0}

    def test_missing_file_raises(self, converter, tmp_path):
        with pytest.raises(FileNotFoundError):
            converter.get_excel_width(file_path=tmp_path / "nope.xlsx")
